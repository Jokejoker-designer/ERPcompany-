# -*- coding: utf-8 -*-
"""Batch 6 acceptance gates. Every write uses synthetic in-memory SQLite."""
import hashlib
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document

import api
import api_write as AW
import docgen
from test_batch2_journal_workflow import make_conn, sess


KTT = sess(1, "ktt1", "Ky thuat truong")
KTT2 = sess(4, "ktt2", "Ky thuat truong")
KTV = sess(2, "ktv1", "Ky thuat vien")
DIRECTOR = sess(31, "gd1", "Giam doc")
ACCOUNTANT = sess(32, "kt1", "Ke toan")


class AcceptanceFixture(unittest.TestCase):
    def setUp(self):
        self.conn = make_conn()
        self.temp_dir = tempfile.TemporaryDirectory()
        for uid, username, role in ((31, "gd1", "Giam doc"), (32, "kt1", "Ke toan")):
            self.conn.execute("""INSERT INTO app_user
                (id,username,full_name,password_hash,salt,role,active,must_change)
                VALUES(?,?,?,?,?,?,1,0)""", (uid, username, username, "hash", "salt", role))
        _, self.source50_hash = self._add_source(
            1, "acceptance-evidence.pdf", b"acceptance-evidence", source_id=50)
        self.conn.execute("""INSERT INTO nhat_ky_thi_cong
            (id,project_id,ngay_ghi,noi_dung,boq_stage_qty_id,khoi_luong_thuc_hien,
             khong_su_dung_vat_tu,created_by,trang_thai,confirmed_by)
            VALUES(100,1,'2026-07-14','Synthetic approved log',1,5,1,2,'Da_duyet',1)""")
        self.conn.commit()

    def tearDown(self):
        self.conn.close()
        self.temp_dir.cleanup()

    def _add_source(self, project_id, name, content, source_id=None):
        path = Path(self.temp_dir.name) / name
        path.write_bytes(content)
        digest = hashlib.sha256(content).hexdigest()
        columns = "customer_id,project_id,profile_role,doc_type,file_name,rel_path,abs_path,ext,size_bytes,source_sha256"
        values = (1, project_id, "attachment", "Ho so", name, "fixture/" + name,
                  str(path), Path(name).suffix or ".pdf", len(content), digest)
        if source_id is None:
            cur = self.conn.execute("INSERT INTO source_document(%s) VALUES(?,?,?,?,?,?,?,?,?,?)" % columns,
                                    values)
            return cur.lastrowid, digest
        self.conn.execute("INSERT INTO source_document(id,%s) VALUES(?,?,?,?,?,?,?,?,?,?,?)" % columns,
                          (source_id,) + values)
        return source_id, digest

    def mark_prerequisites_ready(self, acceptance_type="Giai_doan"):
        projection = api.ct_acceptance(self.conn, "Ky thuat truong", KTT, 1)
        gate = projection["new_draft"]["dossier_gate"]
        for code in gate["required_codes"]:
            source_id, digest = self._add_source(
                1, "%s.pdf" % code, ("acceptance-gate:" + code).encode("utf-8"))
            if code == "CT-05-NKTC":
                self.conn.execute("""INSERT INTO document_export_artifact
                    (project_id,template_code,record_type,record_id,record_version,
                     source_document_id,source_sha256,output_format,generator_version,active,created_by)
                    VALUES(1,'CT-05-NKTC','nhat_ky_thi_cong',100,1,?,?,'PDF','test-fixture',1,1)""",
                                  (source_id, digest))
            self.conn.execute("""INSERT INTO cong_trinh_ho_so_trang_thai
                (project_id,ma_mau,trang_thai,evidence_source_document_id,version,updated_by)
                VALUES(1,?,'Da_duyet',?,1,1)
                ON CONFLICT(project_id,ma_mau) DO UPDATE SET
                  trang_thai='Da_duyet',evidence_source_document_id=excluded.evidence_source_document_id""",
                              (code, source_id))
        self.conn.commit()

    def create_draft(self, **overrides):
        payload = {"phase": "preview", "project_id": 1,
                   "acceptance_type": "Giai_doan", "scope_stage_id": 1,
                   "period_from": "2026-07-01", "period_to": "2026-07-14",
                   "items": [{"boq_stage_qty_id": 1, "acceptance_qty": 5}]}
        payload.update(overrides)
        preview = AW.ct_acceptance_draft(self.conn, KTT, payload)
        return AW.ct_acceptance_draft(self.conn, KTT, {
            "phase": "commit", "confirm_token": preview["confirm_token"]})


class AcceptanceProjectionTest(AcceptanceFixture):
    def test_projection_is_exact_stage_scoped_and_money_free(self):
        result = api.ct_acceptance(self.conn, "Ky thuat truong", KTT, 1)
        draft = result["new_draft"]
        self.assertEqual("LOCKED_V3_1_PROFILE_TRIGGER", draft["dossier_gate"]["policy"])
        self.assertEqual([1], [row["boq_stage_qty_id"] for row in draft["quantity_rows"]])
        self.assertEqual(5, draft["quantity_rows"][0]["journal_confirmed_qty"])
        forbidden = {"don_gia", "unit_price", "amount", "gia_von", "margin", "cong_no"}
        self.assertFalse(forbidden.intersection(str(result).lower()))
        with self.assertRaises(api.PermissionError):
            api.ct_acceptance(self.conn, "Ky thuat vien", KTV, 1)

    def test_upstream_dossier_gate_excludes_its_own_output_and_future_closeout(self):
        stage = api.ct_acceptance(self.conn, "Ky thuat truong", KTT, 1)["new_draft"]
        codes = set(stage["dossier_gate"]["required_codes"])
        self.assertIn("CT-06-WIR", codes)
        self.assertNotIn("CT-06-BBNTGD", codes)
        self.assertNotIn("CT-07-BBBG", codes)
        self.assertEqual("CT-06-BBNTGD", stage["pack_gate"]["template_code"])


class AcceptanceWriteTest(AcceptanceFixture):
    def test_draft_token_is_user_bound_one_time_and_exact_boq_scoped(self):
        preview = AW.ct_acceptance_draft(self.conn, KTT, {
            "phase": "preview", "project_id": 1, "acceptance_type": "Giai_doan",
            "scope_stage_id": 1,
            "items": [{"boq_stage_qty_id": 1, "acceptance_qty": 5}]})
        with self.assertRaises(AW.ValidationError):
            AW.ct_acceptance_draft(self.conn, KTT2, {
                "phase": "commit", "confirm_token": preview["confirm_token"]})
        committed = AW.ct_acceptance_draft(self.conn, KTT, {
            "phase": "commit", "confirm_token": preview["confirm_token"]})
        self.assertEqual("Draft", committed["status"])
        with self.assertRaises(AW.ValidationError):
            AW.ct_acceptance_draft(self.conn, KTT, {
                "phase": "commit", "confirm_token": preview["confirm_token"]})
        with self.assertRaises(AW.ValidationError):
            AW.ct_acceptance_draft(self.conn, KTT, {
                "phase": "preview", "project_id": 1, "acceptance_type": "Giai_doan",
                "scope_stage_id": 1,
                "items": [{"boq_stage_qty_id": 2, "acceptance_qty": 1}]})

    def test_discrepancy_and_material_cocq_fail_closed(self):
        with self.assertRaises(AW.ValidationError):
            self.create_draft(items=[{"boq_stage_qty_id": 1, "acceptance_qty": 4}])
        draft = self.create_draft(items=[{"boq_stage_qty_id": 1, "acceptance_qty": 4,
            "discrepancy_confirmed": True, "discrepancy_reason": "Đối chiếu hiện trường"}])
        self.conn.execute("UPDATE nhat_ky_thi_cong SET khong_su_dung_vat_tu=0 WHERE id=100")
        self.conn.execute("""INSERT INTO nhat_ky_vat_tu
            (nhat_ky_id,boq_stage_qty_id,ten_vat_tu,so_luong_su_dung)
            VALUES(100,1,'Synthetic material',2)""")
        self.conn.commit()
        row = api.ct_acceptance(self.conn, "Ky thuat truong", KTT, 1,
                                acceptance_id=draft["id"])["acceptance"]
        self.assertFalse(row["material_gate"]["ready"])
        self.assertIn("MATERIAL_TRACE_OR_COCQ_INCOMPLETE", row["blockers"])

    def test_submit_then_independent_director_approval_never_marks_signed(self):
        self.mark_prerequisites_ready()
        draft = self.create_draft()
        self.conn.execute("""UPDATE project_acceptance SET
            report_document_id=50,report_sha256=?,report_template_code='CT-06-BBNTGD',
            version=version+1 WHERE id=?""",
                          (self.source50_hash, draft["id"]))
        self.conn.execute("""INSERT INTO cong_trinh_ho_so_trang_thai
            (project_id,ma_mau,trang_thai,evidence_source_document_id,version,updated_by)
            VALUES(1,'CT-06-BBNTGD','Dang_soan',50,1,1)""")
        self.conn.commit()
        current = api.ct_acceptance(self.conn, "Ky thuat truong", KTT, 1,
                                    acceptance_id=draft["id"])["acceptance"]
        submit_preview = AW.ct_acceptance_submit(self.conn, KTT, {
            "phase": "preview", "id": draft["id"], "expected_version": current["version"]})
        submitted = AW.ct_acceptance_submit(self.conn, KTT, {
            "phase": "commit", "confirm_token": submit_preview["confirm_token"]})
        with self.assertRaises(AW.WritePermissionError):
            AW.ct_acceptance_decide(self.conn, KTT, {
                "phase": "preview", "id": draft["id"], "decision": "approve",
                "expected_version": submitted["version"]})
        decision = AW.ct_acceptance_decide(self.conn, DIRECTOR, {
            "phase": "preview", "id": draft["id"], "decision": "approve",
            "expected_version": submitted["version"]})
        approved = AW.ct_acceptance_decide(self.conn, DIRECTOR, {
            "phase": "commit", "confirm_token": decision["confirm_token"]})
        self.assertEqual("Da_duyet", approved["status"])
        self.assertEqual("Chua_ky", approved["signature_status"])
        self.assertIsNone(approved.get("signed_at"))

    def test_pack_preview_commit_uses_exact_v31_template_and_atomic_index(self):
        self.mark_prerequisites_ready()
        draft = self.create_draft()
        with tempfile.TemporaryDirectory() as tmp:
            captured = {}

            def save_and_index(conn, customer_id, _doc_type, filename, data,
                               project_id=None, profile_role=None, commit=True):
                self.assertFalse(commit)
                path = os.path.join(tmp, filename)
                Path(path).write_bytes(data)
                conn.execute("""INSERT INTO source_document
                    (customer_id,project_id,profile_role,doc_type,file_name,rel_path,abs_path,ext,source_sha256)
                    VALUES(?,?,?,?,?,?,?,?,?)""",
                    (customer_id, project_id, profile_role, "Ho so", filename,
                     os.path.join("fixture", filename), path, ".docx",
                     hashlib.sha256(data).hexdigest()))
                captured["bytes"] = data
                return {"ok": True, "abs_path": path}

            preview = AW.ct_acceptance_pack(self.conn, KTT, {
                "phase": "preview", "id": draft["id"], "expected_version": draft["version"]})
            self.assertEqual("CT-06-BBNTGD", preview["template_code"])
            with mock.patch.object(docgen, "luu_file_vao_folder_khach", side_effect=save_and_index):
                committed = AW.ct_acceptance_pack(self.conn, KTT, {
                    "phase": "commit", "confirm_token": preview["confirm_token"]})
            self.assertEqual("Chua_ky", committed["signature_status"])
            self.assertEqual(hashlib.sha256(captured["bytes"]).hexdigest().upper(),
                             committed["report_sha256"])
            row = self.conn.execute("SELECT * FROM project_acceptance WHERE id=?", (draft["id"],)).fetchone()
            self.assertEqual("CT-06-BBNTGD", row["report_template_code"])
            self.assertIsNotNone(row["report_document_id"])
            dossier = self.conn.execute("""SELECT * FROM cong_trinh_ho_so_trang_thai
                WHERE project_id=1 AND ma_mau='CT-06-BBNTGD'""").fetchone()
            self.assertEqual("Dang_soan", dossier["trang_thai"])
            self.assertEqual(row["report_document_id"], dossier["evidence_source_document_id"])


class AcceptanceTemplateTest(unittest.TestCase):
    def test_exact_v31_bbnt_templates_render_dynamic_rows_without_blank_placeholders(self):
        rows = [{"stt": 1, "stage_item": "Tầng 1 — Lắp ống gió", "period_from": "01/07/2026",
                 "period_to": "14/07/2026", "planned_display": "10 m",
                 "acceptance_display": "5 m", "evidence": "Nhật ký đã duyệt",
                 "result": "Đạt", "conclusion": "Đạt"}]
        for code in ("CT-06-BBNTGD", "CT-06-BBNTHH"):
            info = docgen.ct_templates()[code]
            specs = docgen._ct_acceptance_dynamic_spec(code, rows, [])
            output = docgen._export_docx(info["abs_path"], {"TEN_CONG_TRINH": "Fixture"}, [], specs)
            doc = Document(__import__("io").BytesIO(output))
            text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            self.assertIn("Tầng 1 — Lắp ống gió", text)
            self.assertNotIn("{{...}}", text)
            self.assertGreater(len(output), 30000)
            self.assertEqual(64, len(hashlib.sha256(output).hexdigest()))


class AcceptanceFrontendContractTest(unittest.TestCase):
    def test_project_acceptance_workspace_is_wired(self):
        text = Path("web/app_write.js").read_text(encoding="utf-8")
        for marker in ("nghiem_thu", "ct_acceptance", "ct_acceptance_draft",
                       "ct_acceptance_submit", "ct_acceptance_decide", "ct_acceptance_pack"):
            self.assertIn(marker, text)


if __name__ == "__main__":
    unittest.main()
